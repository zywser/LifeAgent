"""笔记加载与切分。"""
from pathlib import Path
from pypdf import PdfReader
from langchain_text_splitters import RecursiveCharacterTextSplitter

_splitter = RecursiveCharacterTextSplitter(
    chunk_size=800,
    chunk_overlap=100,
    separators=["\n\n", "\n", "。", "！", "？", ".", "!", "?", " ", ""],
)


def load_text(filename: str, data: bytes) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix in {".md", ".txt"}:
        return data.decode("utf-8", errors="ignore")
    if suffix == ".pdf":
        import io
        reader = PdfReader(io.BytesIO(data))
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    if suffix in {".docx", ".doc"}:
        import io
        from docx import Document
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(parts)
    raise ValueError("only .md, .txt, .pdf and .docx are supported")


def split_text(text: str) -> list[str]:
    chunks = _splitter.split_text(text)
    return chunks or [""]
